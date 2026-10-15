from docx import Document
from docx.shared import Pt

document = Document()

# Title
title = document.add_heading("Employee Handbook", 0)
document.add_paragraph("Fun Digital Solutions")
document.add_paragraph(
    "This handbook is a fictional document created for an academic "
    "AI Employee Handbook Assistant with RAG project."
)

# 1. Introduction
document.add_heading("1. Introduction", level=1)
document.add_paragraph(
    "Welcome to Fun Digital Solutions. This is a employee handbook "
    "explains important workplace policies, employee responsibilities, "
    "benefits, safety requirements, and company procedures."
)

# 2. Working Hours
document.add_heading("2. Working Hours", level=1)
document.add_paragraph(
    "The standard working time is 37.5 hours per week, normally 7.5 hours "
    "per day from Monday to Friday. Core working hours are from 10:00 to "
    "15:00. Employees should agree on flexible working arrangements with "
    "their manager. Overtime must be approved by the manager in advance."
)

# 3. Annual Leave
document.add_heading("3. Annual Leave", level=1)
document.add_paragraph(
    "Employees are entitled to 25 working days of paid annual leave per year. "
    "New employees receive annual leave on a pro-rata basis. Annual leave "
    "requests should normally be submitted at least two weeks in advance "
    "and require manager approval."
)

# 4. Sick Leave
document.add_heading("4. Sick Leave", level=1)
document.add_paragraph(
    "Employees who are unable to work because of illness must inform their "
    "manager as soon as possible. Employees may use self-certification for "
    "up to three consecutive calendar days. An absence longer than three "
    "days requires a medical certificate or other documentation accepted by HR."
)

# 5. Remote Work
document.add_heading("5. Remote Work", level=1)
document.add_paragraph(
    "Employees may work remotely up to three days per week when their job "
    "is suitable for remote work and their manager has approved the arrangement. "
    "Working remotely from another country requires prior written approval "
    "from both HR and the employee's manager."
)

# 6. Workplace Conduct
document.add_heading("6. Workplace Conduct", level=1)
document.add_paragraph(
    "Employees must behave respectfully and professionally. Harassment, "
    "discrimination, bullying, threats, and abusive behaviour are not permitted. "
    "Employees may report workplace concerns in good faith without retaliation."
)

# 7. Travel and Expenses
document.add_heading("7. Travel and Expenses", level=1)
document.add_paragraph(
    "Business travel must be approved by a manager before booking. "
    "Reasonable business travel costs may be reimbursed. Expense claims "
    "should be submitted within 30 days and receipts should be provided "
    "when available. Personal purchases are not reimbursable."
)

# 8. Health and Safety
document.add_heading("8. Health and Safety", level=1)
document.add_paragraph(
    "The company is committed to maintaining a safe workplace. Employees "
    "must report hazards, accidents, and unsafe conditions. Employees should "
    "follow emergency instructions and safety procedures. Remote workers "
    "should maintain a safe and ergonomic workspace."
)

# 9. IT and Data Security
document.add_heading("9. IT and Data Security", level=1)
document.add_paragraph(
    "Employees must use strong passwords and multi-factor authentication "
    "where required. Passwords must not be shared. Employees must use approved "
    "company systems, lock unattended computers, and report phishing attempts, "
    "lost devices, malware, or suspected data breaches immediately. "
    "Unauthorized software must not be installed on company devices."
)

# 10. Contact and Support
document.add_heading("10. Contact and Support", level=1)
document.add_paragraph(
    "For HR questions, employees may contact hr@fundigitalsolution-example.com. "
    "For technical support, employees may contact it-support@fundigitalsolution-example.com. "
    "For urgent work-related matters, employees should contact their direct manager."
)

# Save the document
document.save("data/employee_handbook.docx")

print("Employee handbook created successfully!")
